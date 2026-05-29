from flask import Flask, request, jsonify
from flask_cors import CORS
from openpyxl import Workbook, load_workbook
from datetime import datetime
import os

app = Flask(__name__)
CORS(app)

EXCEL_FILE = "requests.xlsx"

HEADERS = [
    "ID",
    "Title",
    "Request Type",
    "Department",
    "Priority",
    "Description",
    "Status",
    "Submitted Date"
]


def create_excel_file():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Clinical Requests"
    sheet.append(HEADERS)
    workbook.save(EXCEL_FILE)


def ensure_excel_file_exists():
    if not os.path.exists(EXCEL_FILE):
        create_excel_file()


@app.route("/requests", methods=["POST"])
def create_request():
    ensure_excel_file_exists()

    data = request.get_json()

    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook.active

    request_id = int(datetime.now().timestamp())

    new_row = [
        request_id,
        data.get("title"),
        data.get("requestType"),
        data.get("department"),
        data.get("priority"),
        data.get("description"),
        "No Action",
        datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    ]

    sheet.append(new_row)
    workbook.save(EXCEL_FILE)

    return jsonify({
        "message": "Request saved to Excel successfully",
        "request": {
            "id": request_id,
            "title": data.get("title"),
            "requestType": data.get("requestType"),
            "department": data.get("department"),
            "priority": data.get("priority"),
            "description": data.get("description"),
            "status": "No Action"
        }
    }), 201


@app.route("/requests", methods=["GET"])
def get_requests():
    ensure_excel_file_exists()

    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook.active

    requests = []

    for row in sheet.iter_rows(min_row=2, values_only=True):
        requests.append({
            "id": row[0],
            "title": row[1],
            "requestType": row[2],
            "department": row[3],
            "priority": row[4],
            "description": row[5],
            "status": row[6],
            "submittedDate": row[7]
        })

    return jsonify(requests)


@app.route("/requests/<int:request_id>", methods=["PATCH"])
def update_request_status(request_id):
    ensure_excel_file_exists()

    data = request.get_json()
    new_status = data.get("status")

    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook.active

    for row_number in range(2, sheet.max_row + 1):
        current_id = sheet.cell(row=row_number, column=1).value

        if current_id == request_id:
            sheet.cell(row=row_number, column=7).value = new_status
            workbook.save(EXCEL_FILE)

            return jsonify({
                "message": "Request status updated in Excel successfully"
            })

    return jsonify({"error": "Request not found"}), 404

@app.route("/requests/<int:request_id>", methods=["DELETE"])
def delete_request(request_id):
    ensure_excel_file_exists()

    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook.active

    for row_number in range(2, sheet.max_row + 1):
        current_id = sheet.cell(row=row_number, column=1).value

        if current_id == request_id:
            sheet.delete_rows(row_number)
            workbook.save(EXCEL_FILE)

            return jsonify({
                "message": "Request deleted successfully"
            })

    return jsonify({
        "error": "Request not found"
    }), 404

if __name__ == "__main__":
    ensure_excel_file_exists()
    app.run(debug=True)